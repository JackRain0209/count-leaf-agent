#!/usr/bin/env python3
"""Run the main stalk algorithm on the first 15 Excel samples and build HTML."""

from __future__ import annotations

import html
import json
import shutil
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

import cv2
import numpy as np
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont, ImageOps

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.pipeline.analyzer import analyze_plant_image_stream  # noqa: E402


DATA_DIR = ROOT / "online_uploads/农生院卢坤-合川考种照片-整理20260414"
EXCEL_PATH = DATA_DIR / "实验记录.xlsx"
PHOTO_ROOT = DATA_DIR / "合川考种照片整理-20260414"
REPORT_ROOT = ROOT / "reports"
RESULTS_ROOT = ROOT / "results"
CACHE_ROOT = ROOT / "benchmark_cache"
FONT_PATH = Path("/System/Library/Fonts/Hiragino Sans GB.ttc")


def _font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    try:
        return ImageFont.truetype(str(FONT_PATH), size)
    except Exception:
        return ImageFont.load_default()


def _sample_id(row: dict[str, Any]) -> str:
    existing = row.get("样本编号")
    if existing:
        return str(existing)
    n = row["氮肥浓度"]
    rep = row["重复"]
    variety = row["品种编号"]
    photo_no = row["照片编号"]
    return f"{n}.{rep}.{variety}.{photo_no}"


def _find_photo(sample: str, existing: str | None = None) -> Path:
    if existing:
        p = ROOT / existing
        if p.exists():
            return p
    matches = sorted(PHOTO_ROOT.rglob(f"{sample}.JPG"))
    if not matches:
        matches = sorted(PHOTO_ROOT.rglob(f"{sample}.jpg"))
    if not matches:
        raise FileNotFoundError(f"photo not found for {sample}")
    return matches[0]


def _safe_name(text: str) -> str:
    return "".join(ch if ch.isalnum() or ch in "._-" else "_" for ch in text)


def _accuracy(pred: int, gt: int) -> float:
    if not gt:
        return 0.0
    return round(max(0.0, (1.0 - abs(pred - gt) / gt) * 100.0), 2)


def _rel(path: Path, base: Path) -> str:
    return path.resolve().relative_to(base.resolve()).as_posix()


def _resize_for_report(img: Image.Image, max_w: int = 1600, max_h: int = 1100) -> Image.Image:
    img = ImageOps.exif_transpose(img.convert("RGB"))
    img.thumbnail((max_w, max_h), Image.Resampling.LANCZOS)
    return img


def _save_original(photo: Path, out: Path) -> None:
    img = _resize_for_report(Image.open(photo), 1600, 1100)
    img.save(out, quality=90)


def _draw_bbox_image(photo: Path, plants: list[dict[str, Any]], out: Path) -> None:
    img = _resize_for_report(Image.open(photo), 1800, 1300)
    original = ImageOps.exif_transpose(Image.open(photo).convert("RGB"))
    sx = img.width / original.width
    sy = img.height / original.height
    draw = ImageDraw.Draw(img)
    font = _font(max(22, img.width // 48))

    palette = {
        "主干": (90, 170, 255),
        "主枝": (255, 45, 45),
        "分枝": (40, 220, 115),
    }
    for p in plants:
        bbox = p.get("pixel_bbox")
        if not bbox or len(bbox) != 4:
            continue
        label = p.get("label", "?")
        color = palette.get(label, (255, 220, 80))
        x1, y1, x2, y2 = bbox
        box = [x1 * sx, y1 * sy, x2 * sx, y2 * sy]
        width = 5 if label == "主枝" else 3
        draw.rectangle(box, outline=color, width=width)
        text = f"#{p.get('id')} {label}"
        tx, ty = box[0] + 6, max(4, box[1] - font.size - 8)
        text_color = (255, 0, 0) if label == "主枝" else color
        draw.rectangle(
            [tx - 4, ty - 3, tx + draw.textlength(text, font=font) + 6, ty + font.size + 5],
            fill=(0, 0, 0),
        )
        draw.text((tx, ty), text, fill=text_color, font=font)
    img.save(out, quality=90)


def _letterbox(img: Image.Image, size: tuple[int, int], fill=(22, 28, 38)) -> Image.Image:
    target_w, target_h = size
    img = img.convert("RGB")
    img.thumbnail((target_w, target_h), Image.Resampling.LANCZOS)
    canvas = Image.new("RGB", size, fill)
    canvas.paste(img, ((target_w - img.width) // 2, (target_h - img.height) // 2))
    return canvas


def _make_contact_sheet(
    items: list[tuple[str, Path]],
    out: Path,
    title: str,
    tile_size: tuple[int, int] = (340, 230),
) -> None:
    font_title = _font(28)
    font_label = _font(19)
    if not items:
        canvas = Image.new("RGB", (900, 180), (22, 28, 38))
        d = ImageDraw.Draw(canvas)
        d.text((24, 24), title, fill=(240, 245, 255), font=font_title)
        d.text((24, 82), "无可用图片", fill=(180, 190, 205), font=font_label)
        canvas.save(out, quality=90)
        return

    tile_w, tile_h = tile_size
    label_h = 34
    cols = min(3, max(1, len(items)))
    rows = (len(items) + cols - 1) // cols
    margin = 18
    header_h = 48
    w = cols * tile_w + (cols + 1) * margin
    h = header_h + rows * (tile_h + label_h + margin) + margin
    canvas = Image.new("RGB", (w, h), (22, 28, 38))
    d = ImageDraw.Draw(canvas)
    d.text((margin, 12), title, fill=(240, 245, 255), font=font_title)

    for idx, (label, path) in enumerate(items):
        r, c = divmod(idx, cols)
        x = margin + c * (tile_w + margin)
        y = header_h + r * (tile_h + label_h + margin)
        try:
            tile = _letterbox(Image.open(path), (tile_w, tile_h))
        except Exception:
            tile = Image.new("RGB", (tile_w, tile_h), (45, 20, 20))
        canvas.paste(tile, (x, y))
        d.rectangle([x, y, x + tile_w, y + tile_h], outline=(64, 76, 96), width=1)
        d.text((x + 6, y + tile_h + 6), label, fill=(215, 224, 238), font=font_label)

    canvas.save(out, quality=90)


def _event_url_to_path(url: str) -> Path | None:
    if not url or not url.startswith("/results/"):
        return None
    rel = url.split("?", 1)[0][len("/results/"):]
    p = RESULTS_ROOT / rel
    return p if p.exists() else None


def _final_annotation_items(events: list[dict[str, Any]]) -> list[tuple[str, Path]]:
    items: list[tuple[str, Path]] = []
    for ev in events:
        desc = ev.get("description", "")
        if "VLM 角果复核" not in desc:
            continue
        p = _event_url_to_path(ev.get("url", ""))
        if p:
            label = desc.split("]", 1)[0].replace("[", "").strip() or f"step {ev.get('step')}"
            items.append((label, p))
    return items


def _crop_items(sample: str, plants: list[dict[str, Any]]) -> list[tuple[str, Path]]:
    cache = CACHE_ROOT / sample
    items: list[tuple[str, Path]] = []
    for p in plants:
        if p.get("label") == "主干":
            continue
        pid = p.get("id")
        label = p.get("label", "?")
        matches = sorted(cache.glob(f"crop_{pid}_*.jpg"))
        if matches:
            items.append((f"#{pid} {label} 预测{p.get('pod_count', '-')}", matches[0]))
    return items


def _build_report(report_dir: Path, batch_id: str, rows: list[dict[str, Any]]) -> Path:
    def _value(value: Any) -> str:
        if value is None:
            return ""
        return html.escape(str(value))

    def _percent(value: Any) -> str:
        if value is None or value == "":
            return ""
        return f"{float(value):.2f}%"

    css = """
    body{margin:0;background:#101722;color:#e6eef8;font-family:-apple-system,BlinkMacSystemFont,"Segoe UI",sans-serif}
    main{max-width:1440px;margin:0 auto;padding:28px}
    h1{font-size:28px;margin:0 0 6px}
    .muted{color:#9fb0c6}
    table{border-collapse:collapse;width:100%;margin:20px 0;background:#172234}
    th,td{border:1px solid #304157;padding:8px 10px;text-align:right}
    th:first-child,td:first-child,th:nth-child(2),td:nth-child(2){text-align:left}
    th{background:#203047;color:#cfe2ff}
    section{border:1px solid #304157;border-radius:10px;margin:22px 0;padding:18px;background:#172234}
    h2{margin:0 0 12px;font-size:22px}
    .grid{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:14px}
    figure{margin:0;background:#101722;border:1px solid #2d3d54;border-radius:8px;padding:10px}
    figcaption{font-weight:700;color:#cfe2ff;margin:0 0 8px}
    img{max-width:100%;height:auto;display:block;border-radius:6px}
    .stats{display:flex;gap:12px;flex-wrap:wrap;margin-bottom:12px}
    .pill{background:#223149;border:1px solid #3a4c66;border-radius:999px;padding:5px 10px;color:#d8e6f7}
    """
    html_rows = []
    for r in rows:
        html_rows.append(
            "<tr>"
            f"<td>{html.escape(r['sample'])}</td>"
            f"<td>{html.escape(r['status'])}</td>"
            f"<td>{_value(r.get('gt_branch_count'))}</td><td>{_value(r.get('pred_branch_count'))}</td>"
            f"<td>{_percent(r.get('acc_branch_count'))}</td>"
            f"<td>{r['gt_main']}</td><td>{r['pred_main']}</td>"
            f"<td>{_percent(r.get('acc_main', _accuracy(r['pred_main'], r['gt_main'])))}</td>"
            f"<td>{r['gt_branch']}</td><td>{r['pred_branch']}</td>"
            f"<td>{r['gt_total']}</td><td>{r['pred_total']}</td>"
            f"<td>{r['err_total']:+d}</td><td>{r['acc_total']:.2f}%</td>"
            f"<td>{_value(r.get('main_stem_length', ''))}</td>"
            f"<td>{r['vlm_fp']}</td><td>{r['vlm_restore']}</td>"
            "</tr>"
        )

    sections = []
    for r in rows:
        assets = r["assets"]
        figures = [
            ("1. 原始图", assets["original"]),
            ("2. 分支框线图（主枝红字）", assets["bbox"]),
            ("3. 分支图", assets["branches"]),
            ("4. 分支最终标注图", assets["final"]),
        ]
        fig_html = "\n".join(
            f"<figure><figcaption>{cap}</figcaption><img src='{html.escape(path)}'></figure>"
            for cap, path in figures
        )
        sections.append(
            f"<section id='{html.escape(r['sample'])}'>"
            f"<h2>{html.escape(r['sample'])}</h2>"
            f"<div class='stats'>"
            f"<span class='pill'>真实分支数 {_value(r.get('gt_branch_count'))}</span>"
            f"<span class='pill'>预测分支数 {_value(r.get('pred_branch_count'))}</span>"
            f"<span class='pill'>分支数准确率 {_percent(r.get('acc_branch_count'))}</span>"
            f"<span class='pill'>主花序准确率 {_percent(r.get('acc_main', _accuracy(r['pred_main'], r['gt_main'])))}</span>"
            f"<span class='pill'>主干长度 {_value(r.get('main_stem_length', ''))}</span>"
            f"<span class='pill'>真实总数 {r['gt_total']}</span>"
            f"<span class='pill'>预测总数 {r['pred_total']}</span>"
            f"<span class='pill'>误差 {r['err_total']:+d}</span>"
            f"<span class='pill'>准确率 {r['acc_total']:.2f}%</span>"
            f"<span class='pill'>VLM剔除 {r['vlm_fp']}</span>"
            f"<span class='pill'>VLM补回 {r['vlm_restore']}</span>"
            f"</div><div class='grid'>{fig_html}</div></section>"
        )

    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    page = f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>stalk 15样本实验报告 {batch_id}</title>
<style>{css}</style></head><body><main>
<h1>stalk 主算法 15 样本实验报告</h1>
<div class="muted">批次：{batch_id} ｜ 生成时间：{now} ｜ 样本数：{len(rows)}</div>
<table>
<thead><tr><th>样本</th><th>状态</th><th>真实分支数</th><th>预测分支数</th><th>分支数准确率</th>
<th>真实主花序</th><th>预测主花序</th><th>主花序准确率</th>
<th>真实分枝</th><th>预测分枝</th><th>真实总数</th><th>预测总数</th>
<th>总误差</th><th>总准确率</th><th>主干长度</th><th>VLM剔除</th><th>VLM补回</th></tr></thead>
<tbody>{''.join(html_rows)}</tbody></table>
{''.join(sections)}
</main></body></html>"""
    out = report_dir / "index.html"
    out.write_text(page, encoding="utf-8")
    return out


def main() -> None:
    batch_id = datetime.now().strftime("%Y%m%d_%H%M%S")
    report_dir = REPORT_ROOT / f"stalk_15_report_{batch_id}"
    assets_dir = report_dir / "assets"
    report_dir.mkdir(parents=True, exist_ok=True)
    assets_dir.mkdir(parents=True, exist_ok=True)

    backup = EXCEL_PATH.with_name(
        f"实验记录.backup_before_stalk_15_report_{batch_id}.xlsx")
    shutil.copy2(EXCEL_PATH, backup)
    print(f"[Setup] batch_id={batch_id}")
    print(f"[Setup] excel_backup={backup}")
    print(f"[Setup] report_dir={report_dir}")

    wb = load_workbook(EXCEL_PATH)
    ws = wb.active
    headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}

    def cell(row: int, name: str):
        return ws.cell(row, headers[name])

    rows_summary: list[dict[str, Any]] = []
    target_rows = list(range(2, 17))

    for index, row_idx in enumerate(target_rows, 1):
        row_data = {name: cell(row_idx, name).value for name in headers}
        sample = _sample_id(row_data)
        photo = _find_photo(sample, row_data.get("照片路径"))
        rel_photo = _rel(photo, ROOT)
        run_id = f"{sample}__{batch_id}"
        sample_assets = assets_dir / _safe_name(sample)
        sample_assets.mkdir(parents=True, exist_ok=True)

        print(f"[{index}/15] START row={row_idx} sample={sample} photo={rel_photo}", flush=True)
        start_time = datetime.now()
        events: list[dict[str, Any]] = []
        result: dict[str, Any] | None = None
        error = ""

        try:
            for event in analyze_plant_image_stream(
                str(photo),
                method="stalk",
                run_id=run_id,
                cache_key=sample,
                source_label=rel_photo,
            ):
                if event.get("type") == "step":
                    events.append(event)
                    print(f"  step {event.get('step')}: {event.get('description')}", flush=True)
                elif event.get("type") == "result":
                    result = event
                elif event.get("type") == "error":
                    error = event.get("message", "unknown error")
                    print(f"  ERROR: {error}", flush=True)
            if result is None and not error:
                error = "no result returned"
        except Exception as exc:
            error = str(exc)
            print(f"  EXCEPTION: {error}", flush=True)

        end_time = datetime.now()
        plants = result.get("plants", []) if result else []
        pred_main = sum(int(p.get("pod_count", 0)) for p in plants if p.get("label") == "主枝")
        pred_branch = sum(int(p.get("pod_count", 0)) for p in plants if p.get("label") == "分枝")
        pred_total = int(result.get("total_pods", 0)) if result else 0
        pred_branch_count = len([p for p in plants if p.get("label") == "分枝"])
        pred_main_count = len([p for p in plants if p.get("label") == "主枝"])
        gt_branch_count = int(row_data.get("分枝数") or 0)
        gt_main = int(row_data.get("主花序角果数") or 0)
        gt_branch = int(row_data.get("分枝角果数") or 0)
        gt_total = int(row_data.get("单株总角果数") or 0)
        acc_branch_count = _accuracy(pred_branch_count, gt_branch_count)
        acc_main = _accuracy(pred_main, gt_main)
        main_stem_length = ""
        vlm_fp = sum(len((p.get("vlm_verify") or {}).get("false_positives", [])) for p in plants)
        vlm_restore = sum(len((p.get("vlm_verify") or {}).get("restored_filtered", [])) for p in plants)
        status = "done" if result and not error else "error"
        elapsed = round((end_time - start_time).total_seconds(), 2)

        cell(row_idx, "实验状态").value = status
        cell(row_idx, "实验方法").value = "stalk"
        cell(row_idx, "实验批次").value = batch_id
        cell(row_idx, "样本编号").value = sample
        cell(row_idx, "处理开始时间").value = start_time.strftime("%Y-%m-%d %H:%M:%S")
        cell(row_idx, "处理结束时间").value = end_time.strftime("%Y-%m-%d %H:%M:%S")
        cell(row_idx, "处理耗时秒").value = elapsed
        cell(row_idx, "照片路径").value = rel_photo
        cell(row_idx, "缓存目录").value = _rel(CACHE_ROOT / sample, ROOT)
        cell(row_idx, "结果目录").value = _rel(Path(result.get("result_dir", RESULTS_ROOT / run_id)), ROOT) if result else ""
        cell(row_idx, "识别分支数").value = pred_branch_count
        cell(row_idx, "识别分支数准确率").value = acc_branch_count
        cell(row_idx, "识别主花序角果数").value = pred_main
        cell(row_idx, "识别主花序角果数准确率").value = acc_main
        cell(row_idx, "识别分枝角果数").value = pred_branch
        cell(row_idx, "识别分枝角果数准确率").value = _accuracy(pred_branch, gt_branch)
        cell(row_idx, "识别总角果数").value = pred_total
        cell(row_idx, "识别总角果数准确率").value = _accuracy(pred_total, gt_total)
        cell(row_idx, "识别主枝数").value = pred_main_count
        cell(row_idx, "识别分枝部件数").value = pred_branch_count
        cell(row_idx, "VLM部件数").value = len([p for p in plants if p.get("label") != "主干"])
        cell(row_idx, "VLM误判剔除数").value = vlm_fp
        cell(row_idx, "错误信息").value = error
        wb.save(EXCEL_PATH)

        artifact_dir = CACHE_ROOT / sample
        artifact_dir.mkdir(parents=True, exist_ok=True)
        with (artifact_dir / "experiment_result.json").open("w", encoding="utf-8") as f:
            json.dump({
                "sample": sample,
                "photo_path": rel_photo,
                "method": "stalk",
                "batch_id": batch_id,
                "run_id": run_id,
                "created_at": end_time.isoformat(timespec="seconds"),
                "error": error,
                "result": result,
                "summary": {
                    "gt_branch_count": gt_branch_count,
                    "pred_branch_count": pred_branch_count,
                    "acc_branch_count": acc_branch_count,
                    "gt_main": gt_main,
                    "pred_main": pred_main,
                    "acc_main": acc_main,
                    "gt_branch": gt_branch,
                    "pred_branch": pred_branch,
                    "gt_total": gt_total,
                    "pred_total": pred_total,
                    "main_stem_length": main_stem_length,
                    "vlm_fp": vlm_fp,
                    "vlm_restore": vlm_restore,
                },
            }, f, ensure_ascii=False, indent=2)

        original_out = sample_assets / "01_original.jpg"
        bbox_out = sample_assets / "02_branch_boxes.jpg"
        crops_out = sample_assets / "03_branch_crops.jpg"
        final_out = sample_assets / "04_final_annotations.jpg"
        if result:
            _save_original(photo, original_out)
            _draw_bbox_image(photo, plants, bbox_out)
            _make_contact_sheet(_crop_items(sample, plants), crops_out, "分支图")
            _make_contact_sheet(
                _final_annotation_items(events),
                final_out,
                "分支最终标注图",
                tile_size=(560, 390),
            )
        else:
            _save_original(photo, original_out)
            shutil.copy2(original_out, bbox_out)
            _make_contact_sheet([], crops_out, "分支图")
            _make_contact_sheet([], final_out, "分支最终标注图")

        rows_summary.append({
            "sample": sample,
            "status": status,
            "gt_branch_count": gt_branch_count,
            "pred_branch_count": pred_branch_count,
            "acc_branch_count": acc_branch_count,
            "gt_main": gt_main,
            "pred_main": pred_main,
            "acc_main": acc_main,
            "gt_branch": gt_branch,
            "pred_branch": pred_branch,
            "gt_total": gt_total,
            "pred_total": pred_total,
            "err_total": pred_total - gt_total,
            "acc_total": _accuracy(pred_total, gt_total),
            "main_stem_length": main_stem_length,
            "vlm_fp": vlm_fp,
            "vlm_restore": vlm_restore,
            "assets": {
                "original": _rel(original_out, report_dir),
                "bbox": _rel(bbox_out, report_dir),
                "branches": _rel(crops_out, report_dir),
                "final": _rel(final_out, report_dir),
            },
        })

        print(
            f"[{index}/15] DONE {sample} status={status} "
            f"main={pred_main}/{gt_main} branch={pred_branch}/{gt_branch} "
            f"total={pred_total}/{gt_total} fp={vlm_fp} restore={vlm_restore} "
            f"elapsed={elapsed}s",
            flush=True,
        )

    report = _build_report(report_dir, batch_id, rows_summary)
    summary_path = report_dir / "summary.json"
    summary_path.write_text(json.dumps(rows_summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"[Summary] batch_id={batch_id} processed={len(rows_summary)} failed={sum(1 for r in rows_summary if r['status']!='done')}")
    print(f"[Summary] excel={EXCEL_PATH}")
    print(f"[Summary] backup={backup}")
    print(f"[Summary] report={report}")


if __name__ == "__main__":
    main()
