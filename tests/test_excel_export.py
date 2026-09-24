import io
import tempfile
import unittest
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

import app.main as main_module
from app.history import HistoryStore


# ---------- 辅助函数 ----------

def _display_name(label: str) -> str:
    """导出时图片名会去掉扩展名，断言时用这个函数转换。"""
    return Path(label).stem


def _workbook_text(content: bytes) -> str:
    """把所有 sheet 的单元格文本拼成一个字符串，便于断言。"""
    wb = load_workbook(io.BytesIO(content))
    parts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell is not None:
                    parts.append(str(cell))
    return "\n".join(parts)


def _total_rows(content: bytes) -> int:
    wb = load_workbook(io.BytesIO(content))
    return sum(ws.max_row for ws in wb.worksheets)


# ---------- 测试类 ----------

class ExcelExportTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.original_store = main_module.history_store
        self.original_result_dir = main_module.RESULT_DIR
        self.original_upload_dir = main_module.UPLOAD_DIR

        main_module.history_store = HistoryStore(self.root / "history.sqlite3")
        main_module.RESULT_DIR = self.root / "results"
        main_module.RESULT_DIR.mkdir()
        main_module.UPLOAD_DIR = self.root / "uploads"
        main_module.UPLOAD_DIR.mkdir()
        self.client = TestClient(main_module.app)

        store = main_module.history_store

        # 1) completed，含"枯枝"
        store.start("run-zhizhi-ok", "枯枝正例1.png", "uploads/a.jpg")
        store.complete("run-zhizhi-ok", {
            "type": "result", "run_id": "run-zhizhi-ok",
            "plant_count": 1,
            "plants": [{"id": 1, "label": "主干", "pod_count": 0}],
            "summary": {},
        }, [])

        # 2) completed，含"反例"
        store.start("run-fanli-ok", "枯枝反例1.png", "uploads/b.jpg")
        store.complete("run-fanli-ok", {
            "type": "result", "run_id": "run-fanli-ok",
            "plant_count": 1,
            "plants": [{"id": 1, "label": "主干", "pod_count": 0}],
            "summary": {},
        }, [])

        # 3) failed
        store.start("run-zhizhi-fail", "枯枝反例2.png", "uploads/c.jpg")
        store.fail("run-zhizhi-fail", "VLM 未返回结构", [])

    def tearDown(self):
        self.client.close()
        main_module.history_store = self.original_store
        main_module.RESULT_DIR = self.original_result_dir
        main_module.UPLOAD_DIR = self.original_upload_dir
        self.temp_dir.cleanup()

    # ---------- 基础 ----------

    def test_export_empty_params_returns_valid_xlsx(self):
        """空参数导出：应返回 200 且生成合法的 xlsx 文件。"""
        resp = self.client.get("/api/history/export")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheet", resp.headers["content-type"])
        self.assertGreater(len(resp.content), 0)
        wb = load_workbook(io.BytesIO(resp.content))
        self.assertGreaterEqual(len(wb.sheetnames), 1)

    def test_export_run_ids_only_target_records(self):
        """传 run_ids 时，只导出指定的记录。"""
        resp = self.client.get(
            "/api/history/export",
            params={"run_ids": "run-zhizhi-ok"},
        )
        self.assertEqual(resp.status_code, 200)
        text = _workbook_text(resp.content)
        self.assertIn(_display_name("枯枝正例1.png"), text)
        self.assertNotIn(_display_name("枯枝反例1.png"), text)

    def test_export_run_ids_overrides_q(self):
        """同时传 run_ids 和 q 时，忽略 q，只导出 run_ids。"""
        resp = self.client.get(
            "/api/history/export",
            params={"q": "反例", "run_ids": "run-zhizhi-ok"},
        )
        self.assertEqual(resp.status_code, 200)
        text = _workbook_text(resp.content)
        self.assertIn(_display_name("枯枝正例1.png"), text)
        self.assertNotIn(_display_name("枯枝反例1.png"), text)

    def test_export_duplicate_run_ids_deduped(self):
        """重复的 run_id 应去重，不产生重复行。"""
        single = self.client.get(
            "/api/history/export", params={"run_ids": "run-zhizhi-ok"}
        )
        triple = self.client.get(
            "/api/history/export",
            params={"run_ids": "run-zhizhi-ok,run-zhizhi-ok,run-zhizhi-ok"},
        )
        self.assertEqual(single.status_code, 200)
        self.assertEqual(triple.status_code, 200)
        self.assertEqual(_total_rows(single.content), _total_rows(triple.content))

    # ---------- q 筛选 ----------

    def test_export_by_q_filters(self):
        """按 q 筛选，只导出匹配的记录。"""
        resp = self.client.get("/api/history/export", params={"q": "枯枝"})
        self.assertEqual(resp.status_code, 200)
        text = _workbook_text(resp.content)
        self.assertIn(_display_name("枯枝正例1.png"), text)
        self.assertIn(_display_name("枯枝反例1.png"), text)

    def test_export_q_no_match_returns_empty(self):
        """q 无匹配时，应返回空结果，不崩溃。"""
        resp = self.client.get(
            "/api/history/export",
            params={"q": "这个词肯定没有zzz"},
        )
        self.assertEqual(resp.status_code, 200)
        text = _workbook_text(resp.content)
        self.assertNotIn(_display_name("枯枝正例1.png"), text)
        self.assertNotIn(_display_name("枯枝反例1.png"), text)

    # ---------- 无效输入 ----------

    def test_export_invalid_run_id_does_not_crash(self):
        """传入不存在的 run_id 时，不应崩溃，文件仍可打开。"""
        resp = self.client.get(
            "/api/history/export",
            params={"run_ids": "no-such-run-id"},
        )
        self.assertEqual(resp.status_code, 200)
        load_workbook(io.BytesIO(resp.content))

    def test_export_mixed_valid_and_invalid_run_ids(self):
        """混合有效和无效 run_id 时，只导出有效的那条。"""
        resp = self.client.get(
            "/api/history/export",
            params={"run_ids": "run-zhizhi-ok,no-such-run-id"},
        )
        self.assertEqual(resp.status_code, 200)
        text = _workbook_text(resp.content)
        self.assertIn(_display_name("枯枝正例1.png"), text)


if __name__ == "__main__":
    unittest.main()